from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from xml.sax.saxutils import escape

OUT = Path(__file__).parent
pages = [
[
('name', 'ALEXANDRE AMBIEHL'),
('title', 'Expert technique — Architecture applicative, Java & IA générative'),
('intro', 'Développement logiciel • Architecture applicative • RAG & agents IA'),
('heading', 'PROFIL'),
('text', 'Professionnel de l’informatique depuis 2008, avec un parcours en administration systèmes et réseaux, développement full stack et responsabilité R&D. Expert technique chez CGI depuis mai 2024, spécialisé en architecture applicative et développement. Solide pratique de Java, renforcée par une année sur un projet entièrement Java, et expérience confirmée de l’écosystème PHP/Symfony. Perfectionnement actif en IA générative, systèmes RAG et création d’agents IA, avec l’objectif de contribuer au développement et à l’intégration de solutions IA dans les applications métier.'),
('heading', 'EXPÉRIENCE RÉCENTE'),
('role', 'CGI | Expert technique — Architecture applicative et développement'),
('date', 'Mai 2024 – aujourd’hui'),
('bullet', 'Intervention en tant qu’expert technique en architecture applicative et développement.'),
('bullet', 'Développement sur un projet entièrement Java depuis un an ; consolidation d’une maîtrise avancée du langage par une pratique quotidienne en projet.'),
('bullet', 'Travaux de développement avec Drupal.'),
('text', 'Technologies mises en pratique : Java, Drupal.'),
('heading', 'IA GÉNÉRATIVE — PRATIQUE ET PERFECTIONNEMENT'),
('bullet', 'Utilisation de Codex, Claude, OpenCode et GitHub Copilot pour le développement assisté par IA ; perfectionnement continu des pratiques.'),
('bullet', 'Pratique et approfondissement des systèmes RAG (Retrieval-Augmented Generation, ou génération augmentée par récupération d’informations).'),
('bullet', 'Création d’agents IA et approfondissement de leurs usages.'),
('bullet', 'Connaissance de Symfony AI et de Symfony 8.'),
('heading', 'COMPÉTENCES CLÉS'),
('text', 'Développement : Java, PHP, Symfony (versions 4 à 8), Drupal, Laravel, JavaScript, TypeScript, Vue.js, Svelte/SvelteKit, HTML/CSS, Tailwind CSS.'),
('text', 'Architecture et données : architecture applicative, API, DDD, ADR, PostgreSQL, MySQL, MongoDB, Doctrine, Prisma, RabbitMQ, MQTT.'),
('text', 'Infrastructure et qualité : maîtrise de macOS ; Linux, Windows Server, Docker, CI/CD, GitLab, GitHub, Drone, Nginx, AWS, Azure, PHPUnit.'),
('text', 'Pilotage : cadrage des besoins, cahiers des charges, coordination technique, accompagnement des équipes, Agile, Scrum, Kanban, Jira.'),
],
[
('name', 'ALEXANDRE AMBIEHL'),
('intro', 'Architecture applicative • Développement • Innovation'),
('heading', 'EXPÉRIENCES ANTÉRIEURES'),
('role', 'SERVIALUX | Responsable R&D'),
('date', 'Septembre 2017 – novembre 2023 | Dudelange, Luxembourg'),
('bullet', 'Recueil des besoins clients, rédaction de cahiers des charges et coordination des travaux techniques, des essais utilisateurs et de la documentation.'),
('bullet', 'Conception et développement d’applications web et d’API avec Symfony, PHP, TypeScript, Svelte et Vue.js.'),
('bullet', 'Conception de cartes électroniques et développement de firmwares C++ pour des objets connectés à base d’ESP32.'),
('bullet', 'Intégration de communications entre équipements et services avec RabbitMQ ; gestion des données avec PostgreSQL, MySQL et MongoDB.'),
('bullet', 'Mise en place de pipelines CI/CD avec Docker, Drone et GitLab ; tests PHPUnit et support technique.'),
('text', 'Environnement : Symfony, PHP, TypeScript, Svelte, Vue.js, Tailwind CSS, Doctrine, Prisma, Docker, RabbitMQ, AWS, Azure, C++, PlatformIO, ESP32.'),
('role', 'DISTRISERV | Développeur web full stack'),
('date', 'Janvier 2012 – août 2017'),
('bullet', 'Conception et développement de sites et d’applications web avec PHP, Laravel et Symfony ; intégration des interfaces et des bases de données.'),
('bullet', 'Développement de solutions IoT sur Raspberry Pi avec C++ et Python pour la collecte de données de capteurs et l’automatisation.'),
('bullet', 'Participation à la planification, aux revues de code et aux pratiques agiles au sein d’équipes de développement.'),
('bullet', 'Réalisation de tests, correction d’anomalies et amélioration des performances et de la stabilité des applications.'),
('text', 'Environnement : PHP, Laravel, Symfony, Twig, HTML, CSS, JavaScript, MySQL, Raspberry Pi, C++, Python.'),
('role', 'SERVIATEC | Administrateur réseau'),
('date', 'Août 2008 – décembre 2011 | Metz, France'),
('bullet', 'Déploiement et administration de serveurs et de réseaux sous Linux et Windows Server.'),
('bullet', 'Support technique et assistance aux utilisateurs par hotline.'),
('heading', 'FORMATION'),
('text', '2006 – 2008 | BEP Vente — Anne de Méjanès'),
('text', 'Parcours informatique autodidacte et apprentissage continu en développement, architecture et intelligence artificielle.'),
('heading', 'LANGUES'),
('text', 'Français • Anglais'),
]]

styles = {}
specs = {'name': (23, 27, '#142D42'), 'title': (13, 17, '#176B82'), 'intro': (10, 14, '#536775'), 'heading': (11, 15, '#176B82'), 'role': (11, 15, '#142D42'), 'date': (9, 12, '#536775'), 'text': (10, 14, '#233847'), 'bullet': (10, 14, '#233847')}
for key, (size, leading, color) in specs.items():
    styles[key] = ParagraphStyle(key, fontName='Helvetica-Bold' if key in ['name','title','heading','role'] else 'Helvetica', fontSize=size, leading=leading, textColor=colors.HexColor(color), spaceBefore=9 if key in ['heading','role'] else 0, spaceAfter=5, leftIndent=9 if key == 'bullet' else 0)
story = []
doc = Document()
sec = doc.sections[0]
sec.header_distance = Inches(.3)
sec.top_margin = sec.bottom_margin = Inches(.55)
sec.left_margin = sec.right_margin = Inches(.65)
sec.page_width = Inches(8.27)
sec.page_height = Inches(11.69)
doc.styles['Normal'].font.name = 'Calibri'
doc.styles['Normal'].font.size = Pt(10)
doc.styles['Normal'].paragraph_format.space_after = Pt(5)
md = []
for index, page in enumerate(pages):
    if index:
        story.append(PageBreak())
        doc.add_page_break()
    for kind, content in page:
        story.append(Paragraph(('• ' if kind == 'bullet' else '') + escape(content), styles[kind]))
        para = doc.add_paragraph(style='List Bullet' if kind == 'bullet' else None)
        run = para.add_run(content)
        size, _, color = specs[kind]
        run.font.size = Pt(size)
        run.font.color.rgb = RGBColor.from_string(color[1:])
        run.bold = kind in ['name', 'title', 'heading', 'role']
        if kind in ['heading','role']:
            para.paragraph_format.space_before = Pt(9)
            para.paragraph_format.keep_with_next = True
        md.append(('# ' if kind == 'name' else '## ' if kind == 'heading' else '- ' if kind == 'bullet' else '') + content)
SimpleDocTemplate(str(OUT/'CV_Alexandre_Ambiehl_Java_IA.pdf'), pagesize=(595.28,841.89), rightMargin=40, leftMargin=40, topMargin=32, bottomMargin=32, title='Alexandre Ambiehl — Expert technique Java & IA générative', author='Alexandre Ambiehl').build(story)
doc.save(OUT/'CV_Alexandre_Ambiehl_Java_IA.docx')
(OUT/'CV_Alexandre_Ambiehl_Java_IA.md').write_text('\n\n'.join(md)+'\n')
from pypdf import PdfReader
pdf = PdfReader(OUT/'CV_Alexandre_Ambiehl_Java_IA.pdf')
print(f'PDF : {len(pdf.pages)} pages')
for i, page in enumerate(pdf.pages):
    print(f'Page {i+1} : {len(page.extract_text())} caractères')
